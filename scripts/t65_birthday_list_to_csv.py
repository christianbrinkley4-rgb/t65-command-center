#!/usr/bin/env python3
"""Convert a "T65 <N> NC <Month> <Year> Birthdays" list into the import CSV.

    python scripts/t65_birthday_list_to_csv.py OUT.csv "T65 2092 NC May 1962 Birthdays.xlsx"

A fifth intake shape, and the cleanest one so far: a real header row, one
person per row, no pagination junk. Two things still bite.

**HomePhone is not a phone column.** It holds one of three things: the literal
"DNC", a real ten-digit second number, or nothing. In the May file that's 1,018
DNC flags, 641 real numbers, and 142 blanks in the same column. Reading it as a
phone imports "DNC" as a number; reading it as a flag throws away 641 real
lines, 157 of which are the ONLY number that person has. So it's split by
inspection: DNC sets the compliance flag, ten digits becomes a phone, and which
phone it becomes depends on whether PhoneT was populated.

**The birthday is a placeholder.** Every row in the May file reads 5/1/1962 —
the list is "everyone turning 65 in May", not a list of real birth dates. The
month is true and the day is not, which is fine because everything in the app
files by T65 month, but do not let anyone build a birthday-card feature on it.

Couples share an address and sometimes a landline, so two rows can be two real
people; the importer's samePerson() check is what keeps them apart.
"""

import csv
import re
import sys

import openpyxl

try:
    from oscr_xlsx_to_csv import normalize_city
except ImportError:  # run from the repo root
    sys.path.insert(0, "scripts")
    from oscr_xlsx_to_csv import normalize_city

FIELDS = [
    "Name", "Primary phone", "Secondary phone", "Phone status", "Street",
    "City", "State", "Zip code", "County", "Birthday", "Lead source", "Notes",
]

MONTH_NAMES = ["january", "february", "march", "april", "may", "june",
               "july", "august", "september", "october", "november", "december"]
MONTHS = {m: i + 1 for i, m in enumerate(MONTH_NAMES)}

# The vendor is not consistent about how it writes the month. April and May
# arrived spelled out ("T65 2092 NC May 1962 Birthdays"); June arrived clipped
# ("T65 2092 Jun 1962 Birthdays_NC"). Matching only full names read June as no
# month at all, which does not fail loudly — it labels all 2,203 rows
# "Imported list" instead of "T65 June", so the month has no source to filter
# on and the verifier calls every row unlabeled. Abbreviations are matched too,
# longest first so "sept" is not answered by "sep".
MONTH_PATTERNS = sorted(
    list(MONTHS.items())
    + [(m[:3], i + 1) for i, m in enumerate(MONTH_NAMES)]
    + [("sept", 9)],
    key=lambda kv: -len(kv[0]),
)


def text(v):
    return "" if v is None else str(v).strip()


def digits(v):
    d = re.sub(r"\D", "", text(v))
    if len(d) == 11 and d.startswith("1"):
        d = d[1:]
    return d if len(d) == 10 else ""


def month_year_from_name(filename):
    """"T65 2092 NC May 1962 Birthdays" -> (5, 1962), and "T65 2092 Jun 1962
    Birthdays_NC" -> (6, 1962). The list name is the only trustworthy statement
    of which month this file is."""
    low = filename.lower()
    month = next((n for m, n in MONTH_PATTERNS if re.search(rf"\b{m}\b", low)), None)
    year = re.search(r"\b(19\d{2})\b", low)
    return month, int(year.group(1)) if year else None


def birthday_of(raw, month, year):
    """Prefer a real date in the cell; fall back to the month from the filename.
    Either way the DAY is not to be trusted (see the module docstring)."""
    t = text(raw)
    m = re.match(r"(\d{1,2})/(\d{1,2})/(\d{4})", t)
    if m:
        mm, dd, yy = int(m.group(1)), int(m.group(2)), int(m.group(3))
        return f"{yy:04d}-{mm:02d}-{dd:02d}"
    if t[:4].isdigit() and len(t) >= 10:
        return t[:10]
    if month and year:
        return f"{year:04d}-{month:02d}-01"
    return ""


NEED = {"FirstName", "LastName", "HomeStreet"}


def find_sheet(wb):
    """The sheet holding the list, which is not always the first one.

    The February file arrives as three sheets: an empty `Sheet2` first, the
    1,673 real rows on `NC`, then an empty `Sheet1`. Reading worksheet zero
    found nothing, and the whole month silently converted to zero leads. So the
    sheet is chosen by its HEADER, not its position.
    """
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        try:
            header = [text(h) for h in next(rows)]
        except StopIteration:
            continue
        if NEED.issubset({h for h in header if h}):
            return ws, header
    return None, []


def convert(path, out_rows):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws, header = find_sheet(wb)
    if ws is None:
        wb.close()
        raise SystemExit(
            f"{path}: no sheet in this workbook has the list columns "
            f"({', '.join(sorted(NEED))}). Sheets: {', '.join(wb.sheetnames)}"
        )
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    idx = {h: i for i, h in enumerate(header) if h}
    need = NEED

    name = path.replace("\\", "/").split("/")[-1]
    month, year = month_year_from_name(name)
    source = f"T65 {['','January','February','March','April','May','June','July','August','September','October','November','December'][month]}" if month else "Imported list"

    def cell(r, key):
        return text(r[idx[key]]) if key in idx and idx[key] < len(r) else ""

    stats = {"rows": 0, "dnc": 0, "two_numbers": 0, "rescued_by_home": 0, "address_only": 0}
    for r in rows[1:]:
        if not r or not any(c not in (None, "") for c in r):
            continue
        first, last = cell(r, "FirstName"), cell(r, "LastName")
        full = " ".join(p for p in (first, last) if p)
        street = cell(r, "HomeStreet")
        if not full and not street:
            continue

        primary_raw = cell(r, "PhoneT")
        home_raw = cell(r, "HomePhone")
        is_dnc = home_raw.upper() == "DNC"
        home_digits = "" if is_dnc else digits(home_raw)
        primary_digits = digits(primary_raw)

        # Which number can actually be dialed, and which is the spare.
        if primary_digits:
            primary, secondary = primary_raw, (home_raw if home_digits and home_digits != primary_digits else "")
            if secondary:
                stats["two_numbers"] += 1
        elif home_digits:
            primary, secondary = home_raw, ""
            stats["rescued_by_home"] += 1
        else:
            primary, secondary = "", ""
            stats["address_only"] += 1

        if is_dnc:
            stats["dnc"] += 1

        out_rows.append({
            "Name": full,
            "Primary phone": primary,
            "Secondary phone": secondary,
            "Phone status": "DNC" if is_dnc else "",
            "Street": street,
            "City": normalize_city(cell(r, "HomeCity")),
            "State": cell(r, "HomeState") or "NC",
            "Zip code": cell(r, "HomePostalCode")[:5],
            "County": cell(r, "County"),
            "Birthday": birthday_of(cell(r, "Birthday"), month, year),
            "Lead source": source,
            "Notes": "",
        })
        stats["rows"] += 1
    return stats, source


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if len(args) < 2:
        raise SystemExit("Usage: t65_birthday_list_to_csv.py OUT.csv LIST.xlsx [more...]")
    out_path, sources = args[0], args[1:]

    rows = []
    for src in sources:
        stats, label = convert(src, rows)
        print(f"{src.split(chr(92))[-1]:<45} {stats['rows']:>5} leads  list={label}")
        print(f"{'':<45} DNC {stats['dnc']}, two numbers {stats['two_numbers']}, "
              f"rescued by HomePhone {stats['rescued_by_home']}, address-only {stats['address_only']}")

    with open(out_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(rows)
    print(f"\n{len(rows)} leads -> {out_path}")


if __name__ == "__main__":
    main()
